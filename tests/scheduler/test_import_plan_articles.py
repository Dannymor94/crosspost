"""Tests for article rows in content-plan import.

Covers:
1. No type column → all rows import as posts (backward compat).
2. Mixed plan: post + article rows in one file.
3. Article body parses into the same blocks as parse_blocks on the same text.
4. Article without title → row-level error; rest of plan still imports.
5. Article with filename in media column → not treated as image, reported explicitly.
6. Synonym matching: type/title column headers in both languages.
"""

from __future__ import annotations

from crosspost.content.article_blocks import ParagraphBlock, parse_blocks
from crosspost.scheduler.import_plan import (
    ScheduleRule,
    build_plan,
    detect_columns,
    parse_table,
)


def _csv(text: str) -> bytes:
    return text.encode("utf-8")


def _rule() -> ScheduleRule:
    from crosspost.scheduler.import_plan import parse_rule

    return parse_rule("2025-01-01", "10:00")


# ── 1. No type column → all rows import as posts ──────────────────────────────


def test_no_type_column_all_rows_are_posts():
    """Plan with no type column: all rows imported as posts, behaviour unchanged."""
    data = _csv("файл,текст\n1.jpg,Привет мир\n2.jpg,Второй пост")
    rows = parse_table(data, "plan.csv")
    assert len(rows) == 2
    for row in rows:
        assert row.content_type == "post"


def test_no_type_column_build_plan_all_posts():
    data = _csv("файл,текст\n1.jpg,Привет мир")
    rows = parse_table(data, "plan.csv")
    report = build_plan(rows, ["1.jpg"], _rule())
    assert len(report.resolved) == 1
    assert report.resolved[0].content_type == "post"
    assert not report.problems


# ── 2. Mixed plan: post rows and article rows ─────────────────────────────────


def test_mixed_plan_post_and_article_rows():
    """Post and article rows coexist in one plan, each gets the right content_type."""
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["тип", "файл", "текст", "заголовок"])
    ws.append(["пост", "cover.jpg", "Текст поста", ""])
    ws.append(["статья", "", "Первый абзац.\n\nВторой абзац.", "Заголовок статьи"])
    buf = _io.BytesIO()
    wb.save(buf)
    rows = parse_table(buf.getvalue(), "plan.xlsx")
    assert len(rows) == 2
    assert rows[0].content_type == "post"
    assert rows[1].content_type == "article"
    assert rows[1].title == "Заголовок статьи"


def test_mixed_plan_build_report():
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["тип", "файл", "текст", "заголовок"])
    ws.append(["пост", "cover.jpg", "Текст поста", ""])
    ws.append(["статья", "", "Абзац.\n\nhttps://example.com/img.png", "Заголовок"])
    buf = _io.BytesIO()
    wb.save(buf)
    rows = parse_table(buf.getvalue(), "plan.xlsx")
    report = build_plan(rows, ["cover.jpg"], _rule())
    assert not report.problems
    post_row = next(r for r in report.resolved if r.content_type == "post")
    article_row = next(r for r in report.resolved if r.content_type == "article")
    assert post_row.file == "cover.jpg"
    assert article_row.title == "Заголовок"
    assert article_row.blocks is not None


# ── 3. Article body parses into blocks matching parse_blocks ──────────────────


def test_article_blocks_match_parse_blocks():
    """Blocks on ResolvedRow equal parse_blocks output on the same text."""
    body = "Введение.\n\nhttps://example.com/img.jpg\n\nЗаключение."
    # CSV doesn't handle embedded newlines well — use xlsx instead
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["тип", "файл", "текст", "заголовок"])
    ws.append(["статья", "", body, "Заголовок"])
    buf = __import__("io").BytesIO()
    wb.save(buf)
    xlsx_data = buf.getvalue()

    rows = parse_table(xlsx_data, "plan.xlsx")
    assert len(rows) == 1
    report = build_plan(rows, [], _rule())
    assert not report.problems
    rr = report.resolved[0]
    assert rr.content_type == "article"
    assert rr.blocks is not None
    expected = parse_blocks(body)
    assert len(rr.blocks) == len(expected)
    for got, want in zip(rr.blocks, expected, strict=True):
        assert type(got) is type(want)
        if isinstance(want, ParagraphBlock):
            assert got.text == want.text
        else:
            assert got.url == want.url


# ── 4. Article without title → row-level error, rest imports ─────────────────


def test_article_without_title_is_row_error():
    """Article row with no title produces a row-level problem; other rows still import."""
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["тип", "файл", "текст", "заголовок"])
    ws.append(["статья", "", "Тело статьи.", ""])  # no title
    ws.append(["пост", "1.jpg", "Нормальный пост", ""])  # should import fine
    buf = _io.BytesIO()
    wb.save(buf)

    rows = parse_table(buf.getvalue(), "plan.xlsx")
    report = build_plan(rows, ["1.jpg"], _rule())
    assert len(report.problems) == 1
    assert (
        "заголовок" in report.problems[0].reason.lower()
        or "title" in report.problems[0].reason.lower()
    )
    assert report.problems[0].row_number == 1
    # The post row still resolves.
    assert len(report.resolved) == 1
    assert report.resolved[0].content_type == "post"


def test_article_without_title_no_silent_skip():
    """Missing title on article: never a silent skip, never a fallback from body."""
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["тип", "файл", "текст"])
    ws.append(["статья", "", "Тело без заголовка."])
    buf = _io.BytesIO()
    wb.save(buf)

    rows = parse_table(buf.getvalue(), "plan.xlsx")
    report = build_plan(rows, [], _rule())
    assert len(report.resolved) == 0
    assert len(report.problems) == 1
    # The resolved body slice is NOT used as the title.
    assert not any(r.title == "Тело без заголовка." for r in report.resolved)


# ── 5. Article with filename in media column → not treated as image ───────────


def test_article_filename_in_media_column_not_resolved_as_image():
    """For article rows, the file column value is not treated as an image source.

    The filename is reported explicitly, not silently ignored.
    """
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["тип", "файл", "текст", "заголовок"])
    ws.append(["статья", "photo.jpg", "Тело статьи.", "Мой заголовок"])
    buf = _io.BytesIO()
    wb.save(buf)

    rows = parse_table(buf.getvalue(), "plan.xlsx")
    report = build_plan(rows, ["photo.jpg"], _rule(), folder_configured=False)
    # Resolved (not dropped), but media_source is not "upload"/"folder"/"http"
    assert len(report.resolved) == 1
    rr = report.resolved[0]
    # Article images come only from URL blocks in the body, not from the file column.
    assert rr.media_source == ""
    assert rr.file == ""  # filename not used as image source
    # There is a warning/problem reporting the ignored filename.
    noted = report.warnings + report.problems
    assert any("photo.jpg" in (p.file + p.reason) for p in noted)


# ── 6. Synonym matching: type/title in both languages ─────────────────────────


def test_synonym_type_column_recognised_in_russian():
    cols = detect_columns(["файл", "текст", "Тип", "Заголовок"])
    assert "type" in cols
    assert "title" in cols


def test_synonym_type_column_recognised_in_english():
    cols = detect_columns(["file", "text", "type", "title"])
    assert "type" in cols
    assert "title" in cols


def test_synonym_type_column_synonyms():
    """All reasonable type synonyms map to 'type'."""
    for header in ["Тип", "тип контента", "content type", "вид", "kind"]:
        cols = detect_columns(["file", "text", header])
        assert "type" in cols, f"synonym {header!r} not recognised"


def test_synonym_title_column_synonyms():
    """All reasonable title synonyms map to 'title'."""
    for header in ["Заголовок", "название", "title", "heading"]:
        cols = detect_columns(["file", "text", header])
        assert "title" in cols, f"synonym {header!r} not recognised"


def test_empty_type_cell_defaults_to_post():
    """Empty type cell → post (not an error)."""
    data = _csv("тип,файл,текст\n,1.jpg,Привет\nстатья,,Тело")
    rows = parse_table(data, "plan.csv")
    assert rows[0].content_type == "post"  # empty → post


def test_unknown_type_cell_is_row_error():
    """Unrecognised type value → row-level error, not a crash."""
    data = _csv("тип,файл,текст\nвидео,1.jpg,Привет")
    rows = parse_table(data, "plan.csv")
    report = build_plan(rows, ["1.jpg"], _rule())
    assert len(report.problems) == 1
    assert "тип" in report.problems[0].reason.lower() or "type" in report.problems[0].reason.lower()
